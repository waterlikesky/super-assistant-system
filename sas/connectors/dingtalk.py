"""钉钉：用户自己导出 / 整理的聊天记录 csv 或 xlsx。

钉钉没有统一的个人导出格式，这里按列名别名识别：
    时间 / 发送时间 / time      发送人 / 发送者 / sender      内容 / 消息内容 / content
    会话 / 群名称 / conversation（可选，缺省用文件名）
xlsx 需要 openpyxl（pip install '.[xlsx]'），缺包时给出明确提示。
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator

from .base import (
    Connector,
    ParseContext,
    SkipRecord,
    has_any,
    make_chat_event,
    parse_time,
    pick,
    read_csv_header,
    read_csv_rows,
)

TIME = ["时间", "发送时间", "消息时间", "time", "timestamp", "date"]
SENDER = ["发送人", "发送者", "发言人", "sender", "from", "nick"]
CONTENT = ["内容", "消息内容", "content", "message", "text"]
CONVERSATION = ["会话", "会话名称", "群名称", "群聊", "conversation", "chat"]


def _rows_from_xlsx(path: Path) -> tuple[list[str], list[dict]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SkipRecord("读取 xlsx 需要 openpyxl：pip install 'super-assistant-system[xlsx]'，或先另存为 csv") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        rows_iter = wb.active.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(rows_iter, [])]
        rows = [
            {header[i]: ("" if v is None else str(v)).strip() for i, v in enumerate(row) if i < len(header)}
            for row in rows_iter
        ]
    finally:
        wb.close()
    return header, rows


class DingTalkConnector(Connector):
    name = "dingtalk"
    channel = "dingtalk"
    description = "钉钉聊天导出 csv / xlsx（时间、发送人、内容[、会话]）"
    suffixes = (".csv", ".xlsx")

    def sniff(self, path: Path) -> bool:
        suffix = path.suffix.lower()
        if suffix == ".xlsx":
            return "钉钉" in path.name or "dingtalk" in path.name.lower()
        if suffix != ".csv":
            return False
        header = read_csv_header(path)
        return has_any(header, TIME) and has_any(header, SENDER) and has_any(header, CONTENT)

    def parse(self, path: Path, ctx: ParseContext) -> Iterator[dict]:
        _, rows = _rows_from_xlsx(path) if path.suffix.lower() == ".xlsx" else read_csv_rows(path)
        by_conv: dict[str, list[tuple]] = {}
        for index, row in enumerate(rows, start=2):
            where = f"{path.name}:{index}"
            text = pick(row, CONTENT)
            sender = pick(row, SENDER)
            if not text:
                ctx.skip(where, "正文为空")
                continue
            try:
                when = parse_time(pick(row, TIME), ctx.tz)
            except SkipRecord as exc:
                ctx.skip(where, exc.reason)
                continue
            conversation = pick(row, CONVERSATION) or path.stem
            by_conv.setdefault(conversation, []).append((when, sender or "未知", ctx.config.is_me(sender), text))
        for conversation, records in by_conv.items():
            others = {s for _, s, me, _ in records if not me}
            is_group = len(others) > 1
            for when, sender, is_me, text in records:
                yield make_chat_event(
                    ctx=ctx, channel="dingtalk", via="dingtalk_export", path=path,
                    conversation=conversation, sender=sender, is_me=is_me, when=when, text=text,
                    peer=None if is_group else (next(iter(others)) if others else conversation), is_group=is_group,
                )
